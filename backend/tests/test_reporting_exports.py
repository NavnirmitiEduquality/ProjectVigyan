import io
from datetime import date

from openpyxl import load_workbook

from app.main import app
from app.services.reporting.exports import metadata, report_bytes
from fastapi.testclient import TestClient
from tests.test_reporting_attendance import _headers
from tests.test_reporting_domains import domain_fixture
from app.models import SessionEvidence, TeachingSession, User
from app.services.reporting.common import (
    ReportingFilters,
    production_session_query,
)
from app.routers.reporting import _evidence_export_rows, _session_export_rows

client = TestClient(app)


def test_csv_xlsx_pdf_exports_preserve_report_rows():
    rows = [{"attendance_percentage": 83.33, "eligible_record_count": 12}]
    report_metadata = metadata("attendance", "tester", {"school_id": None})
    csv_payload, _, _ = report_bytes("attendance", rows, report_metadata, "csv")
    xlsx_payload, _, _ = report_bytes("attendance", rows, report_metadata, "xlsx")
    pdf_payload, _, _ = report_bytes("attendance", rows, report_metadata, "pdf")
    assert b"attendance_percentage" in csv_payload
    workbook = load_workbook(io.BytesIO(xlsx_payload))
    assert workbook.active["A1"].value == "Project Vigyan"
    assert b"%PDF" == pdf_payload[:4]


def test_export_cannot_cross_school_scope(domain_fixture):
    response = client.get(
        "/api/v1/reporting/exports/sessions",
        params={"school_id": str(domain_fixture["school_two"].id)},
        headers=_headers(domain_fixture["viewer_token"]),
    )
    assert response.status_code == 403


def test_http_csv_export(domain_fixture):
    response = client.get(
        "/api/v1/reporting/exports/delivery",
        params={"file_format": "csv"},
        headers=_headers(domain_fixture["platform_owner_token"]),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert response.headers["content-disposition"].endswith(
        'filename="delivery-report.csv"'
    )
    assert b"teaching_hours" in response.content


def test_http_xlsx_export(domain_fixture):
    response = client.get(
        "/api/v1/reporting/exports/delivery",
        params={"file_format": "xlsx"},
        headers=_headers(domain_fixture["platform_owner_token"]),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"].endswith(
        'filename="delivery-report.xlsx"'
    )

    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.active["A1"].value == "Project Vigyan"


def test_http_pdf_export(domain_fixture):
    response = client.get(
        "/api/v1/reporting/exports/delivery",
        params={"file_format": "pdf"},
        headers=_headers(domain_fixture["platform_owner_token"]),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.headers["content-disposition"].endswith(
        'filename="delivery-report.pdf"'
    )
    assert response.content.startswith(b"%PDF")

def test_export_requires_authentication():
    response = client.get("/api/v1/reporting/exports/delivery")
    assert response.status_code == 401


def test_export_requires_report_export_permission(domain_fixture):
    response = client.get(
        "/api/v1/reporting/exports/delivery",
        headers=_headers(domain_fixture["no_permission_token"]),
    )
    assert response.status_code == 403


def test_export_rejects_unsupported_format(domain_fixture):
    response = client.get(
        "/api/v1/reporting/exports/delivery",
        params={"file_format": "json"},
        headers=_headers(domain_fixture["platform_owner_token"]),
    )
    assert response.status_code == 422


def test_export_metadata_is_consistent_across_formats():
    rows = [{"example_value": 10}]
    report_metadata = metadata(
        "delivery",
        "tester",
        {
            "academic_year_id": "year-123",
            "start_date": "2026-06-01",
            "end_date": "2026-06-30",
            "school_id": None,
        },
    )

    csv_payload, _, _ = report_bytes(
        "delivery", rows, report_metadata, "csv"
    )
    xlsx_payload, _, _ = report_bytes(
        "delivery", rows, report_metadata, "xlsx"
    )
    pdf_payload, _, _ = report_bytes(
        "delivery", rows, report_metadata, "pdf"
    )

    for label in (
        b"project_name",
        b"academic_year_id",
        b"report_type",
        b"reporting_period",
        b"generated_at",
        b"generated_by",
        b"applied_filters",
    ):
        assert label in csv_payload

    workbook = load_workbook(io.BytesIO(xlsx_payload))
    xlsx_metadata = "\n".join(
        str(cell.value)
        for row in workbook.active.iter_rows()
        for cell in row
        if cell.value is not None
    )
    for label in (
        "project_name",
        "academic_year_id",
        "report_type",
        "reporting_period",
        "generated_at",
        "generated_by",
        "applied_filters",
    ):
        assert label in xlsx_metadata

    assert pdf_payload.startswith(b"%PDF")


def test_export_metadata_preserves_applied_filters(domain_fixture):
    fixture = domain_fixture

    para_teacher = (
        fixture["db"]
        .query(User)
        .filter(User.user_code.like("RP-PT-%"))
        .one()
    )

    response = client.get(
        "/api/v1/reporting/exports/sessions",
        params={
            "academic_year_id": str(fixture["year"].id),
            "start_date": "2026-06-01",
            "end_date": "2026-06-30",
            "school_id": str(fixture["school_one"].id),
            "class_division_id": str(fixture["division_one"].id),
            "para_teacher_id": str(para_teacher.id),
            "group_by": "school",
            "file_format": "csv",
        },
        headers=_headers(fixture["platform_owner_token"]),
    )

    assert response.status_code == 200
    content = response.content.decode("utf-8")

    assert "applied_filters" in content
    assert str(fixture["year"].id) in content
    assert str(fixture["school_one"].id) in content
    assert str(fixture["division_one"].id) in content
    assert str(para_teacher.id) in content
    assert "2026-06-01" in content
    assert "2026-06-30" in content
    assert "school" in content


def test_session_export_applies_date_and_academic_year_filters(domain_fixture):
    fixture = domain_fixture
    db = fixture["db"]

    filters = ReportingFilters(
        authorized_school_ids=None,
        academic_year_id=fixture["year"].id,
        start_date=date(2026, 6, 2),
        end_date=date(2026, 6, 4),
    )

    rows = _session_export_rows(db, filters)

    assert rows
    assert all(
        date(2026, 6, 2)
        <= row["session_date"]
        <= date(2026, 6, 4)
        for row in rows
    )
    assert all(row["status"] == "COMPLETED" for row in rows)


def test_evidence_export_matches_filtered_production_sessions(domain_fixture):
    fixture = domain_fixture
    db = fixture["db"]

    filters = ReportingFilters(
        authorized_school_ids=None,
        academic_year_id=fixture["year"].id,
        start_date=__import__("datetime").date(2026, 6, 1),
        end_date=__import__("datetime").date(2026, 6, 3),
        school_id=fixture["school_one"].id,
        class_division_id=fixture["division_one"].id,
    )

    session_rows = _session_export_rows(db, filters)
    evidence_rows = _evidence_export_rows(db, filters)
    filtered_session_ids = {row["session_id"] for row in session_rows}

    assert all(
        row["session_id"] in filtered_session_ids for row in evidence_rows
    )

    production_session_ids = {
        session.id
        for session in production_session_query(db).all()
    }
    assert all(
        row["session_id"] in production_session_ids for row in evidence_rows
    )
    assert all(
        row["school_id"] == fixture["school_one"].id for row in evidence_rows
    )
    assert all(
        {"evidence_id", "session_id", "school_id", "created_at", "photo_id"}
        == set(row.keys())
        for row in evidence_rows
    )

def test_pdf_export_escapes_special_characters_in_metadata():
    rows = []
    report_metadata = metadata(
        "delivery",
        "School A & B <North>",
        {"school_name": "School A & B <North>"},
    )

    payload, media_type, extension = report_bytes(
        "delivery", rows, report_metadata, "pdf"
    )

    assert payload.startswith(b"%PDF")
    assert media_type == "application/pdf"
    assert extension == "pdf"
